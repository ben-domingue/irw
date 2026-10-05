#!/usr/bin/env python3
"""Item text for the nine lyu_2024 tables (10.1038/s41598-024-58060-4).

The cheap case in Step 3.5: the article's own supplement
41598_2024_58060_MOESM1_ESM.docx, Table S1 "Survey Instrument", is a two-column
table of item code and item wording for every item, and its codes are the
headers of the data file MOESM2_ESM.csv, which data/lyu_2024_social_entrepreneurship.py
keeps unchanged as `item`. So each stem is tied to its code by name, not by
position. All 42 codes match exactly; no deviations. Stems are
whitespace-normalised only.

Anchors: Methods says respondents rated on "a five-point Likert scale, from
'strongly disagree' (1) to 'strongly agree' (5)"; those two phrases are used for
1 and 5 and points 2-4 are left blank, not padded. No instructions are published.
The administration language is not stated in the paper (Chinese university
sample, online survey); Table S1 is in English and is shipped as the study's own
item list, without language/_translated columns.
"""
import csv
import io
import zipfile
from pathlib import Path

import docx
import pandas as pd
import requests

HERE = Path(__file__).resolve().parent.parent
OUT_DIR = HERE / "itemtext_output"
RESP_DIR = HERE / "irw_output"
NAME = "41598_2024_58060_MOESM1_ESM.docx"
RAW = HERE / "runs" / "raw" / "lyu_2024_social_entrepreneurship" / NAME
UA = {"User-Agent": "IRW-Finder/1.0 (ben.domingue@gmail.com)"}
SUPP = "https://www.ebi.ac.uk/europepmc/webservices/rest/PMC10978997/supplementaryFiles"

COLS = ["table", "section_id", "item", "instrument", "instructions",
        "section_prompt", "item_text", "correct_response", "option_text",
        "resp"]

TABLES = {
    "lyu_2024_rtp": ("Risk-taking propensity", "RTP"),
    "lyu_2024_sef": ("Self-efficacy", "SEF"),
    "lyu_2024_nfa": ("Need for achievement", "NFA"),
    "lyu_2024_pvs": ("Perceived values on sustainability", "PVS"),
    "lyu_2024_org": ("Opportunity recognition competency", "ORG"),
    "lyu_2024_ate": ("Attitude towards entrepreneurship", "ATE"),
    "lyu_2024_sun": ("Subjective norms", "SUN"),
    "lyu_2024_pbc": ("Perceived behavioural control", "PBC"),
    "lyu_2024_sei": ("Social entrepreneurial intention", "SEI"),
}
OPTIONS = {1: "strongly disagree", 2: "", 3: "", 4: "", 5: "strongly agree"}


def s1() -> dict:
    if RAW.exists():
        blob = RAW.read_bytes()
    else:
        r = requests.get(SUPP, headers=UA, timeout=300)
        r.raise_for_status()
        blob = zipfile.ZipFile(io.BytesIO(r.content)).read(NAME)
    doc = docx.Document(io.BytesIO(blob))
    assert len(doc.tables) == 1
    t = doc.tables[0]
    assert [c.text.strip() for c in t.rows[0].cells] == ["Code", "Questions"]
    out = {}
    for row in t.rows[1:]:
        code, text = (c.text.strip() for c in row.cells)
        assert code not in out, code
        out[code] = " ".join(text.split())
    return out


def main() -> None:
    stems = s1()
    assert len(stems) == 42, len(stems)
    used = set()
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for table, (instrument, prefix) in TABLES.items():
        resp = pd.read_csv(RESP_DIR / f"{table}.csv")
        items = sorted(resp["item"].unique(), key=lambda s: int(s[len(prefix):]))
        assert all(i.startswith(prefix) for i in items)
        rows = []
        for it in items:
            stem = stems[it]  # KeyError = no tie by code
            used.add(it)
            for r, opt in OPTIONS.items():
                rows.append({"table": table, "section_id": f"{table}_1", "item": it,
                             "instrument": instrument, "instructions": "",
                             "section_prompt": "", "item_text": stem,
                             "correct_response": "", "option_text": opt, "resp": r})
        assert {r["item"] for r in rows} == set(resp["item"])
        assert {r["resp"] for r in rows} == set(resp["resp"])
        path = OUT_DIR / f"{table}__items.csv"
        with open(path, "w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=COLS, quoting=csv.QUOTE_ALL)
            w.writeheader()
            w.writerows(rows)
        print(f"{path.name}: {len(rows)} rows, {len(items)} items")
    assert used == set(stems), set(stems) - used


if __name__ == "__main__":
    main()
