#!/usr/bin/env python3
"""Item text for prasetiyo_2026_act / _sct / _digcompedu (Mendeley 10.17632/v8k4k2km69).

Source: the deposited "Research Instrument ... .docx". ACT and SCT are numbered tables (No. | Question |
1-4 anchors); the printed number k is the data code ACTk / SCTk (positional by number, nothing else ties
them). DigCompEdu CheckIn: a three-column table (area | item | five ordered option statements, one per
line) in block order PE DR TL AS EL FLDC = the data's PE1-4 DR1-3 TL1-4 AS1-3 EL1-3 FLDC1-4. The
instrument is deposited in English; the administration language is not stated (Indonesian lecturers).
"""
import re

import docx

from _itemtext_writer_c import load_script, row, write

mod = load_script("prasetiyo_2026_digital_competence")
DOC = mod.mendeley_fetch("v8k4k2km69", "Research Instrument Measurement of Digital Competence in Indonesia Higher Education.docx", mod.RAW_DIR)
AGREE4 = {1: "Strongly Disagree", 2: "Slightly Disagree", 3: "Slightly Agree", 4: "Strongly Agree"}
BLOCKS = [("PE", 4), ("DR", 3), ("TL", 4), ("AS", 3), ("EL", 3), ("FLDC", 4)]


def numbered(table):
    out = {}
    for r in table.rows:
        cells = [c.text.strip() for c in r.cells]
        if re.fullmatch(r"\d+", cells[0]):
            out[int(cells[0])] = " ".join(cells[1].split())
    return out


def main():
    d = docx.Document(DOC)
    act, sct, dce = d.tables[2], d.tables[3], d.tables[4]
    act_items, sct_items = numbered(act), numbered(sct)
    assert len(act_items) == 12 and len(sct_items) == 25
    for table, items, prefix, instr in [("prasetiyo_2026_act", act_items, "ACT", "Attitude toward Computer Technology (ACT), Milbrath & Kinzie (2000), 12 items"),
                                        ("prasetiyo_2026_sct", sct_items, "SCT", "Self-Efficacy with Computer Technology (SCT), Milbrath & Kinzie (2000), 25 items")]:
        rows = [row(table, f"{table}_1", f"{prefix}{k}", instr, "", r, stem, option_text=AGREE4[r])
                for k, stem in items.items() for r in AGREE4]
        write(table, rows)
    dce_rows = []
    codes = [f"{p}{i}" for p, n in BLOCKS for i in range(1, n + 1)]
    k = 0
    for r in dce.rows[1:]:
        cells = [c.text.strip() for c in r.cells]
        stem = " ".join(cells[1].split())
        opts = [" ".join(o.split()) for o in cells[2].split("\n") if o.strip()]
        if not stem:
            if opts and dce_rows:  # continuation row: the previous item's remaining option line(s)
                prev = [x for x in dce_rows if x["item"] == dce_rows[-1]["item"]]
                blanks = [x for x in prev if not x["option_text"]]
                for x, o in zip(blanks, opts):
                    x["option_text"] = o
            continue
        code = codes[k]
        k += 1
        if len(opts) != 5:
            print(f"  {code}: {len(opts)} option lines in the row (the rest come from a continuation row)")
            opts = (opts + [""] * 5)[:5]
        for j in range(1, 6):
            dce_rows.append(row("prasetiyo_2026_digcompedu", "prasetiyo_2026_digcompedu_1", code,
                                "DigCompEdu CheckIn self-assessment, 21 items (Horváth et al. 2024 version)", "", j, stem, option_text=opts[j - 1]))
    assert k == 21, k
    write("prasetiyo_2026_digcompedu", dce_rows)


if __name__ == "__main__":
    main()
