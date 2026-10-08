#!/usr/bin/env python3
"""Item text for dassean_2025_bsmas (figshare 10.6084/m9.figshare.28375283.v2).

The deposit's "Supplemantry material.docx" prints the administered Arabic BSMAS as a
table: a header row with the five options and the lead-in "كم مرة خلال العام
الماضي............" ("How often during the last year..."), then one row per item whose
text ends with its item number (1-6). data/dassean_2025_bsmas.py keeps the .sav
names item1-item6, whose SPSS variable labels name the BSMAS components in the
standard order (Salience, Tolerance, Mood modification, Relapse, Withdrawal,
Conflict) -- which the Arabic items 1-6 match in content. So item k = the row
numbered k. The trailing number is stripped from item_text; nothing else changed.

Options: the .sav value labels are English (1 Very rarely ... 5 Very often); the
docx header lists the Arabic options right-to-left from كثيرا جدا (very often) to
نادرا جدا (very rarely), so 1 = نادرا جدا ... 5 = كثيرا جدا; the English value labels
are shipped as option_text_translated (study-supplied). Instructions = the docx's
"التعليمات" paragraph followed by the table lead-in.

item_text_translated and instructions_translated are IRW's own translations
(translation_source = machine_translation in the provenance note; owes an
itemtext_issues.qmd entry).
"""
import csv
import re
from pathlib import Path

import docx
import pandas as pd
import pyreadstat

HERE = Path(__file__).resolve().parent.parent
RAW = HERE / "runs" / "raw" / "dassean_2025"
OUT = HERE / "itemtext_output" / "dassean_2025_bsmas__items.csv"
RESP = HERE / "irw_output" / "dassean_2025_bsmas.csv"
TABLE = "dassean_2025_bsmas"
COLS = ["table", "section_id", "item", "instrument", "language", "instructions",
        "instructions_translated", "section_prompt", "item_text", "item_text_translated",
        "correct_response", "option_text", "option_text_translated", "resp"]
EN = {1: "Spent a lot of time thinking about social media or planning to use it?",
      2: "Felt a constant urge to use social media?",
      3: "Used social media to forget about your personal problems?",
      4: "Tried to cut down on your use of social media without success?",
      5: "Felt anxious or troubled if you were prevented from using social media?",
      6: "Used social media so excessively that it had a negative effect on your work/studies?"}
INSTR_EN = ("Instructions: Below you will find some questions about your relationship with "
            "social media and your use of it (Facebook, YouTube, Twitter, Instagram, "
            "Snapchat, TikTok ........). Choose the alternative answer to each question "
            "that best describes you. How often during the last year............")


def main():
    doc = docx.Document(str(RAW / "supp.docx"))
    (tab,) = doc.tables
    rows = [[c.text.strip() for c in r.cells] for r in tab.rows]
    opts_rtl, lead = rows[0][:5], rows[0][5]
    assert opts_rtl == ["كثيرا جدا", "كثيرا", "احيانا", "نادرا", "نادرا جدا"], opts_rtl
    ar_opts = dict(zip([5, 4, 3, 2, 1], opts_rtl))
    (instr,) = [p.text for p in doc.paragraphs if p.text.strip().startswith("التعليمات")]
    instr = " ".join((instr + " " + lead).split())
    stems = {}
    for r in rows[1:]:
        m = re.fullmatch(r"(.*?)(\d)", r[5], re.S)
        stems[int(m.group(2))] = " ".join(m.group(1).split())
    assert sorted(stems) == [1, 2, 3, 4, 5, 6]
    _, meta = pyreadstat.read_sav(str(RAW / "BSMAS_jordanian_version.sav"), metadataonly=True)
    labels = [meta.column_names_to_labels[f"item{k}"] for k in range(1, 7)]
    assert labels == ["Salience", "Tolerance", "Mood modification", "Relapse",
                      "Withdrawal", "Conflict"], labels
    en_opts = meta.variable_value_labels["item1"]
    resp = pd.read_csv(RESP)
    out = []
    for k in range(1, 7):
        for r in range(1, 6):
            out.append({"table": TABLE, "section_id": f"{TABLE}_1", "item": f"item{k}",
                        "instrument": "Bergen Social Media Addiction Scale (BSMAS), Jordanian version",
                        "language": "Arabic", "instructions": instr,
                        "instructions_translated": INSTR_EN, "section_prompt": "",
                        "item_text": stems[k], "item_text_translated": EN[k],
                        "correct_response": "", "option_text": ar_opts[r],
                        "option_text_translated": en_opts[float(r)], "resp": r})
    assert {x["item"] for x in out} == set(resp["item"]) and set(resp["resp"]) <= {1, 2, 3, 4, 5}
    with open(OUT, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=COLS, quoting=csv.QUOTE_ALL)
        w.writeheader()
        w.writerows(out)
    print(f"{OUT.name}: {len(out)} rows")


if __name__ == "__main__":
    main()
