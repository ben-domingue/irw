#!/usr/bin/env python3
"""Item text for elshafie_2020_edsc (10.7717/peerj.10301).

Administered wording: the deposit's "Arabic checklist" (PeerJ SI
peerj-08-10301-s002.docx), one table, 54 rows. Each row is numbered "k-" and
carries its age band (1 month ... 25-30 months) and the two options
"نعم / لا" (Yes / No). The children's caregivers were interviewed with this
form (Methods: the Baroda items were "translated ... to Arabic", then
"clarified to the parents/or guardians"). The response file's item codes are
the .sav columns Q1..Q54 (value labels 0 = No, 1 = Yes); Qk is tied to the
row numbered k (paper_order), and the age-band ordering is checked in
verify_elshafie_2020_edsc.R.

item_text is the Arabic cell verbatim with the leading "k-" number (and the
stray "-"/"." that follows it on items 27 and 29) removed and runs of spaces
collapsed. item_text_translated is IRW's own English translation of that
Arabic. It is NOT the English Baroda checklist in s001: the Arabic form was
adapted and departs from it item by item (e.g. Arabic 27 is the pincer grasp
where Baroda's 10th-month list has "Pulls string-secures toys").
The form's heading "بيانات التطور العقلى والحركى" ("Mental and motor
development data") is a form title, not an instruction, and is not shipped.
"""
import csv
import re
import zipfile
from io import BytesIO
from pathlib import Path

import docx
import pandas as pd
import requests

HERE = Path(__file__).resolve().parent.parent
OUT = HERE / "itemtext_output" / "elshafie_2020_edsc__items.csv"
RESP = HERE / "irw_output" / "elshafie_2020_edsc.csv"
SUPP = "https://www.ebi.ac.uk/europepmc/webservices/rest/PMC7666562/supplementaryFiles"
RAW = HERE / "runs" / "raw" / "elshafie_2020" / "supp.zip"
UA = {"User-Agent": "IRW-Finder/1.0 (ben.domingue@gmail.com)"}
TABLE = "elshafie_2020_edsc"
INSTRUMENT = ("Egyptian Developmental Screening Chart (EDSC) checklist: Arabic "
              "adaptation of the Baroda Developmental Screening Test, 54 "
              "motor and mental milestones, birth to 30 months")

EN = {
    1: "Moves his arms and legs to the sides (kicks with his arms and legs)",
    2: "Fixes his eyes on the mother's face (or a light) for a few moments",
    3: "Turns his head to the sides while lying on his stomach",
    4: "Startles or goes quiet on hearing ordinary sounds (e.g. talking, a rattle)",
    5: "Follows the mother with his eyes when she moves, or a toy moved in front of him",
    6: "Moves his eyes to explore the place around him, trying to notice something in front of him.",
    7: "Smiles and coos audibly when spoken to or when sounds are made to him.",
    8: "Is the child able to look around in all directions (up-down, right-left)?",
    9: "Raises his head for a short time (wobbly) when the mother carries him on her shoulder",
    10: "Holds his head up steadily when pulled by the arms to sit or lifted upright from under the armpits.",
    11: "Smiles on seeing his mother; calms down when she comes.",
    12: "When lying on his stomach, is he able to lift his chest off the bed, leaning on his elbows, or lift his body leaning on his arms?",
    13: "When a rattle is put in his hand, does the child look at it, put it in his mouth, shake it or play with it in any other way?",
    14: "Does the child show interest in things around him and try to reach them by stretching out his arm (he need not be able to reach them)?",
    15: "Sits supported by a pillow or the mother's arm with the legs stretched out. May be tested by the doctor.",
    16: "Turns towards different sounds (people, the television or a rattle)",
    17: "Rolls from his back onto his side",
    18: "Does the child play with paper when given it (e.g. putting it in his mouth or crumpling it in his hand to make sounds)?",
    19: "Is wary of strangers (cries, turns away, stares)",
    20: "Starts to sit up on his own from lying down when the parents hold out their hands to him, without help pulling the child.",
    21: "Bangs things together or on the floor to make sounds (e.g. blocks, plates)",
    22: "Sits upright without needing support for several minutes.",
    23: "When the child is given a toy in one hand while holding a toy in the other, does he keep holding both toys?",
    24: "When the mother holds out her hands to the child while he is sitting, does he try to pull himself up to stand without help?",
    25: "Laughs or makes sounds when looking in the mirror",
    26: "Sits without help and can turn around, or starts to crawl trying to reach a distant toy, without losing his balance.",
    27: "Picks up small things (a pea or food) like pincers (with two fingers).",
    28: "When you ask him for a toy, he holds out his hand to give it to you, even if he does not let go of it.",
    29: "Crawls.",
    30: "Does the child shake the rattle on purpose to make sounds?",
    31: "Picks up things smoothly with the tips of his index and middle fingers (e.g. a grain of rice)",
    32: "Sits without help",
    33: "Stands holding on to furniture",
    34: "Knows some things by name: when asked where something is, without pointing to it, does the child look at it?",
    35: "Says a two-syllable word: ba-ba, ma-ma, na-na",
    36: "Understands the word \"no\" (\"don't do it\"), (\"put it back\", \"come here\")",
    37: "Claps, or holds a toy on his left side with his right hand, touches his left arm with his right hand, or puts one leg over the other while lying on his back.",
    38: "Walks a few steps while holding the mother's hands (a baby walker or small furniture does not count).",
    39: "Turns pages to look at more pictures (or only tries to tear the paper)",
    40: "Repeats some words after you or makes sounds similar to them",
    41: "Stands without help for a short time after being helped to stand and the support is removed.",
    42: "Scribbles with crayons",
    43: "Throws the ball away when holding it with both hands.",
    44: "Can the child get from lying down to standing without any outside help?",
    45: "Walks a few steps alone without holding on to anything, even if he cannot stand up by himself (when you help him stand and let go, he can walk)",
    46: "Shakes his head to refuse, points to the cup of water to drink, points to a toy he wants and cries for it to be brought",
    47: "Recognizes his own things (his shoes, his toy, his clothes)",
    48: "Says two clear words and knows their meaning (e.g. \"booh\", \"mam\")",
    49: "Holds someone's hand or the railing while going up and down stairs (climbing on hands and knees or sitting does not count)",
    50: "Points to a thing while saying its name correctly (\"ball\"), even in the child's own way of speaking",
    51: "Speaks in a two-word sentence (e.g. \"want drink\", \"want eat\", \"daddy out\")",
    52: "Names three things (e.g. a picture of a dog, cat, horse, daddy, mummy)",
    53: "Stands on one foot, e.g. when (putting on trousers or shoes)",
    54: "Can go up and down without holding the railing",
}
OPTIONS = [(1, "نعم", "Yes"), (0, "لا", "No")]


def arabic_items():
    if RAW.exists():
        zb = RAW.read_bytes()
    else:
        r = requests.get(SUPP, headers=UA, timeout=300)
        r.raise_for_status()
        zb = r.content
    z = zipfile.ZipFile(BytesIO(zb))
    d = docx.Document(BytesIO(z.read("peerj-08-10301-s002.docx")))
    assert len(d.tables) == 1
    out = {}
    for row in d.tables[0].rows:
        for c in row.cells:
            t = c.text.strip()
            m = re.match(r"^(\d+)\s*-\s*[-.]?\s*(.+)$", t, flags=re.S)
            if m:
                k = int(m.group(1))
                txt = re.sub(r"\s+", " ", m.group(2)).strip()
                assert out.get(k, txt) == txt, k
                out[k] = txt
    assert sorted(out) == list(range(1, 55)), sorted(out)
    return out


def main():
    ar = arabic_items()
    resp = pd.read_csv(RESP)
    assert set(resp["item"]) == {f"Q{k}" for k in range(1, 55)}
    assert set(resp["resp"]) == {0, 1}
    rows = []
    for k in range(1, 55):
        for v, opt_ar, opt_en in OPTIONS:
            rows.append({
                "table": TABLE, "section_id": f"{TABLE}_1", "item": f"Q{k}",
                "instrument": INSTRUMENT, "language": "Arabic",
                "instructions": None, "section_prompt": None,
                "item_text": ar[k], "correct_response": None,
                "option_text": opt_ar, "resp": v,
                "instructions_translated": None,
                "section_prompt_translated": None,
                "item_text_translated": EN[k],
                "option_text_translated": opt_en,
            })
    df = pd.DataFrame(rows)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(OUT, index=False, quoting=csv.QUOTE_ALL, na_rep="NA")
    print(f"{OUT.name}: {len(df)} rows, {df['item'].nunique()} items")


if __name__ == "__main__":
    main()
