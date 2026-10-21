#!/usr/bin/env python3
##Minneapolis ranked-choice: compare a FairVote raw file with the city's own cast vote record (CVR) of the same contest.
##Both files are turned into a multiset of (precinct, rank 1, rank 2, ...) and the two multisets are compared.
##Differences are reported, never adjusted away.
##
##Usage:
##  python3 compare_city_cvr.py --raw Minneapolis_<date>_<Office>.tab --city <city cvr file> [--format csv|xlsx] \
##      [--aliases aliases.csv] [--empty-is-skip] [--show N]
##
##City file layouts:
##  csv   one header row, then precinct, one column per rank, and optionally a column named Count (a ballot weight; 1 if
##        there is none). The Minneapolis CSVs of 2017, 2021 and 2023 look like this.
##  xlsx  the 2013 workbook: sheet "Marked Sheet", header Precinct, 1ST CHOICE, 2ND CHOICE, 3RD CHOICE, COUNT.
##        Needs openpyxl.
##The 2009 sheets (three-letter names, typed write-ins) are not read here.
##
##Normalisation. Both sides go through the same rules, and every string that needed one is listed in the report:
##  precinct   surrounding spaces removed; a date-like value such as 2001-01-04 (a precinct 01-04 that a spreadsheet turned
##             into a date) becomes 01-04
##  name       backslashes and double quotes removed (a nickname in quotes is written in several ways), surrounding spaces removed
##  token      the case-, space- and hyphen-insensitive tokens skipped and undervote count as a skip; overvote as an
##             overvote; writein and uwi as a write-in. With --empty-is-skip an empty city cell also counts as a skip (some
##             2023 files leave a cell empty). Any other lowercase single word on the city side is an error: add it to the
##             map in this script after checking what the word stands for.
##  file       a city file that is not valid UTF-8 is read as cp1252; a byte order mark is removed. Both are reported.
##A row with no precinct and no choices is a totals line, not a ballot: it is reported and left out.
##
##Aliases. The city may spell a name differently from the FairVote file (an accent, a middle name). Names are personal
##data and are not kept in this repository, so the spelling map is a file you supply at run time and keep yourself:
##a CSV with the header city,raw and one row per spelling. Without it, every difference in spelling shows up as a difference.
##
##Exit code: 0 the two multisets are identical (and the same size), 1 they differ, 2 bad input.
##Names stay out of the report unless you ask: the normalisation notes for names show only a count per rule. With --show N
##(N above 0) the report also prints the name strings and up to N differing ballots from each side. They contain names,
##so do not paste that output into public text.
from __future__ import annotations

import argparse
import csv
import io
import re
import sys
from collections import Counter
from pathlib import Path

SKIP, OVER, WI = "~skip", "~over", "~wi"
TOKENS = {"skipped": SKIP, "undervote": SKIP, "overvote": OVER, "writein": WI, "uwi": WI}
CANONICAL = {"skipped", "overvote", "writein"}  # FairVote's own spellings: no note needed
DATE_LIKE = re.compile(r"^\d{4}-(\d{2})-(\d{2})$")
RANK = re.compile(r"^rank(\d+)$", re.I)


class Norm:
    def __init__(self, tokens: dict[str, str], aliases: dict[str, str], notes: dict[str, None], strict: bool):
        self.tokens, self.aliases, self.notes, self.strict, self.memo = tokens, aliases, notes, strict, {}

    def note(self, kind: str, before: str, after: str, rule: str) -> None:
        self.notes[f"{kind} {before!r} -> {after!r} ({rule})"] = None

    def precinct(self, v: str) -> str:
        out = v.strip()
        if out != v:
            self.note("precinct", v, out, "strip")
        m = DATE_LIKE.match(out)
        if m:
            before, out = out, f"{m[1]}-{m[2]}"
            self.note("precinct", before, out, "date-like")
        return out

    def cell(self, v: str) -> str:
        if v in self.memo:
            return self.memo[v]
        out = re.sub(r'[\\"]', "", v)
        if out != v:
            self.note("name", v, out, "quotes removed")
        s = out.strip()
        if s != out:
            self.note("name", out, s, "strip")
        key = re.sub(r"[\s\-_]", "", s).lower()
        if key in self.tokens:
            out = self.tokens[key]
            if s not in CANONICAL:
                self.note("token", v, out, "token map")
        elif self.strict and re.fullmatch(r"[a-z]+", s):
            raise ValueError(f"unmapped token {v!r}: check what it stands for, then add it to TOKENS")
        else:
            out = self.aliases.get(s, s)
            if out != s:
                self.note("name", s, out, "alias")
        self.memo[v] = out
        return out


def read_raw(path: Path) -> tuple[list[tuple[str, list[str]]], int]:
    """(precinct, rank cells) per FairVote row, and the number of rank columns."""
    with open(path, newline="", encoding="utf-8") as f:
        reader = csv.reader(f, delimiter="\t")
        header = next(reader)
        rows = list(reader)
    rank = sorted((int(RANK.match(h).group(1)), i) for i, h in enumerate(header) if RANK.match(h))
    if not rank:
        raise ValueError(f"{path.name}: no rank columns")
    pcol = next((i for i, h in enumerate(header) if h.lower() in ("precinct", "precincts")), None)
    need = max([i for _, i in rank] + ([pcol] if pcol is not None else []))
    for n, r in enumerate(rows, start=2):
        if len(r) <= need:
            raise ValueError(f"{path.name}: line {n} has {len(r)} cells, the header needs {need + 1}")
    return [(r[pcol] if pcol is not None else "", [r[i] for _, i in rank]) for r in rows], len(rank)


def _decode(path: Path, notes: dict[str, None]) -> str:
    b = path.read_bytes()
    try:
        text = b.decode("utf-8-sig")
        if b.startswith(b"\xef\xbb\xbf"):
            notes[f"file {path.name!r} -> utf-8 (byte order mark removed)"] = None
        return text
    except UnicodeDecodeError:
        notes[f"file {path.name!r} -> cp1252 (not valid utf-8)"] = None
        return b.decode("cp1252")


def read_city_csv(path: Path, k: int, notes: dict[str, None]) -> list[tuple[str, list[str], int]]:
    r = csv.reader(io.StringIO(_decode(path, notes), newline=""))
    header = next(r)
    count_col = next((i for i, h in enumerate(header) if h.strip().lower() == "count"), None)
    out = []
    for row in r:
        row = row + [""] * (k + 2 - len(row))
        cells = row[1:1 + k]
        if not row[0] and not any(cells):
            notes[f"totals row {','.join(row)!r} ignored"] = None
            continue
        n = int(row[count_col]) if count_col is not None and row[count_col] != "" else 1
        out.append((row[0], cells, n))
    return out


def read_city_xlsx(path: Path, k: int, notes: dict[str, None]) -> list[tuple[str, list[str], int]]:
    try:
        import openpyxl
    except ImportError:
        raise ValueError("reading an xlsx file needs openpyxl (python3 -m pip install openpyxl)") from None
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:  # openpyxl raises its own error types for a file it cannot open
        raise ValueError(f"{path.name}: cannot open as xlsx ({e.__class__.__name__})") from None
    ws = wb["Marked Sheet"] if "Marked Sheet" in wb.sheetnames else wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    head = [str(c or "").strip().lower() for c in next(rows)[:k + 2]]
    want = ("precinct", "1st choice", "2nd choice", "3rd choice", "count")
    if k != 3 or not all(h.startswith(w) for h, w in zip(head, want)):
        raise ValueError(f"{path.name}: unexpected header {head}")

    def text(v) -> str:
        if v is None:
            return ""
        if isinstance(v, (int, float)) and float(v).is_integer():
            return str(int(v))
        return str(v)

    out = []
    for i, row in enumerate(rows, 2):
        row = list(row) + [None] * (5 - len(row))
        cells = [text(c) for c in row[:4]]
        if not cells[0] and not any(cells[1:]):
            notes[f"totals row {','.join(text(c) for c in row[:5])!r} ignored"] = None
            continue
        if row[4] is None or isinstance(row[4], bool) or not float(row[4]).is_integer():
            raise ValueError(f"{path.name}: row {i} has COUNT {row[4]!r}")
        out.append((cells[0], cells[1:], int(row[4])))
    wb.close()
    return out


def read_aliases(path: Path | None) -> dict[str, str]:
    if path is None:
        return {}
    with open(path, newline="", encoding="utf-8-sig") as f:
        rows = list(csv.DictReader(f))
    out: dict[str, str] = {}
    for r in rows:
        k, v = r["city"].strip(), r["raw"]
        if out.setdefault(k, v) != v:
            raise ValueError(f"alias collision: {k!r} maps to two names")
    return out


def compare(raw_path: Path, city_path: Path, fmt: str, aliases: dict[str, str], empty_is_skip: bool):
    notes: dict[str, None] = {}
    raw_rows, k = read_raw(raw_path)
    city_rows = (read_city_xlsx if fmt == "xlsx" else read_city_csv)(city_path, k, notes)
    fv = Norm(TOKENS, {}, notes, strict=False)
    cv = Norm({**TOKENS, **({"": SKIP} if empty_is_skip else {})}, aliases, notes, strict=True)
    a: Counter = Counter()
    for p, cells in raw_rows:
        a[(fv.precinct(p), *map(fv.cell, cells))] += 1
    b: Counter = Counter()
    for p, cells, n in city_rows:
        b[(cv.precinct(p), *map(cv.cell, cells))] += n
    return a, b, list(notes)


def main(argv: list[str]) -> int:
    ap = argparse.ArgumentParser(description="Compare a FairVote raw file with the city's cast vote record.")
    ap.add_argument("--raw", required=True, type=Path)
    ap.add_argument("--city", required=True, type=Path)
    ap.add_argument("--format", choices=("csv", "xlsx"), default="csv")
    ap.add_argument("--aliases", type=Path)
    ap.add_argument("--empty-is-skip", action="store_true")
    ap.add_argument("--show", type=int, default=0)
    args = ap.parse_args(argv)
    try:
        a, b, notes = compare(args.raw, args.city, args.format, read_aliases(args.aliases), args.empty_is_skip)
    except (OSError, ValueError, KeyError, StopIteration) as e:
        print(f"cannot compare: {e!r}", file=sys.stderr)
        return 2
    only_raw, only_city = a - b, b - a
    n_raw, n_city = sum(a.values()), sum(b.values())
    same = not only_raw and not only_city
    print(f"rows: raw {n_raw}, city {n_city}; only in raw {sum(only_raw.values())}, only in city {sum(only_city.values())}; "
          + ("IDENTICAL" if same else "DIFFERENT"))
    if not args.show:  # name notes hold name strings, so show only a count per rule
        names = Counter(n.rsplit(" (", 1)[1].rstrip(")") for n in notes if n.startswith("name "))
        notes = [n for n in notes if not n.startswith("name ")] + [f"name: {c} string(s), rule: {r} (text hidden, use --show)" for r, c in sorted(names.items())]
    print(f"normalised ({len(notes)}):" if notes else "normalised: nothing needed")
    for n in notes:
        print(f"  {n}")
    for label, extra in (("only in raw", only_raw), ("only in city", only_city)):
        for key, n in list(extra.items())[:args.show]:
            print(f"  {label} x{n}: {key}")
    return 0 if same else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
